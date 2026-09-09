#!/usr/bin/env python3
# -*- coding: utf-8 -*-
#
# MOBILE ROBOTS - FI-UNAM
# EXCEL (.xlsx) REPORT GENERATOR FOR A* RESULTS
#
import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generar_excel_a_star(input_csv="a_star_raw_results.csv", output_xlsx="reporte_A_star.xlsx"):
    if not os.path.exists(input_csv):
        print(f"[ERROR] No se encontró el archivo {input_csv}. Ejecuta primero a_star_tester.py.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados A-Star"

    headers = [
        "Punto Destino X (m)",
        "Punto Destino Y (m)",
        "Uso de Diagonales",
        "Número de Nodos",
        "Tiempo de Ejecución (ms)"
    ]
    ws.append(headers)

    with open(input_csv, mode='r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            gx = float(row["goal_x"])
            gy = float(row["goal_y"])
            diag_str = "Sí" if row["use_diagonals"].lower() == "true" else "No"
            nodes = int(row["path_nodes"])
            t_ms = round(float(row["time_ms"]), 2)

            ws.append([gx, gy, diag_str, nodes, t_ms])

    # Estilos visuales
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Aplicar estilos a encabezados
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Aplicar estilos a filas de datos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = center_align
            cell.border = thin_border

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 16)

    wb.save(output_xlsx)
    print(f"\n[ÉXITO] Archivo Excel generado: {os.path.abspath(output_xlsx)}")
    print(f"Total de rutas registradas: {ws.max_row - 1}")

if __name__ == '__main__':
    generar_excel_a_star()